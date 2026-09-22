"""Render an exam to PDF / DOCX in the layout Tunisian teachers expect.

Layout (when a spec with school context is provided):

    ┌──────────────────────────────┬──────────────────────────────┐
    │ Établissement : …            │ Matière : …                  │
    │ Enseignant(e) : …            │ Niveau : …                   │
    │ Année scolaire : …           │ Durée : … min                │
    └──────────────────────────────┴──────────────────────────────┘
                Devoir de contrôle — 1er trimestre
                          <exam title>
    Nom & Prénom : ……………   Classe : ……   Note : …… / 20

    Exercice N°1 (5 points) …
    …
    ── page break ──
    Corrigé et barème

Arabic papers are rendered right-to-left with the bundled Amiri font. ReportLab
does not implement the Unicode bidi algorithm, so Arabic text is shaped
(``arabic_reshaper``) and reordered (``python-bidi``) line by line after we
wrap it ourselves against the frame width.
"""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional
from xml.sax.saxutils import escape

import arabic_reshaper
from bidi.algorithm import get_display
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from ..curriculum import tunisia as cur
from ..models.schemas import ExamContent, ExamSpec, Exercise

logger = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# Fonts
# --------------------------------------------------------------------------- #

_FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
ARABIC_FONT = "Amiri"
ARABIC_FONT_BOLD = "Amiri-Bold"
_arabic_fonts_ready: Optional[bool] = None


def _ensure_arabic_fonts() -> bool:
    """Register the bundled Amiri family once. Returns False if unavailable."""
    global _arabic_fonts_ready
    if _arabic_fonts_ready is None:
        try:
            pdfmetrics.registerFont(
                TTFont(ARABIC_FONT, str(_FONT_DIR / "Amiri-Regular.ttf"))
            )
            pdfmetrics.registerFont(
                TTFont(ARABIC_FONT_BOLD, str(_FONT_DIR / "Amiri-Bold.ttf"))
            )
            pdfmetrics.registerFontFamily(
                ARABIC_FONT,
                normal=ARABIC_FONT,
                bold=ARABIC_FONT_BOLD,
                italic=ARABIC_FONT,
                boldItalic=ARABIC_FONT_BOLD,
            )
            _arabic_fonts_ready = True
        except Exception:  # pragma: no cover - depends on the deployed filesystem
            logger.warning(
                "Arabic fonts not found in %s; Arabic PDFs will not render correctly",
                _FONT_DIR,
            )
            _arabic_fonts_ready = False
    return _arabic_fonts_ready


# --------------------------------------------------------------------------- #
# Paper metadata derived from the spec
# --------------------------------------------------------------------------- #


@dataclass
class _Paper:
    lang: str = "en"
    rtl: bool = False
    labels: dict[str, str] = field(default_factory=dict)
    heading: Optional[str] = None
    # (label, value) pairs for the two header columns.
    left_rows: list[tuple[str, str]] = field(default_factory=list)
    right_rows: list[tuple[str, str]] = field(default_factory=list)

    @property
    def has_header(self) -> bool:
        return bool(self.left_rows or self.right_rows)

    def t(self, key: str) -> str:
        return self.labels.get(key) or cur.PAPER_LABELS["en"].get(key, key)


def _paper_from_spec(spec: Optional[ExamSpec]) -> _Paper:
    lang = (spec.language if spec else "en") or "en"
    paper = _Paper(lang=lang, rtl=(lang == "ar"), labels=cur.PAPER_LABELS.get(lang, {}))
    if spec is None:
        return paper

    if spec.exam_type and spec.exam_type != "generic":
        heading = cur.label(cur.EXAM_TYPE_LABELS, lang, spec.exam_type)
        if spec.trimester:
            heading = f"{heading} — {cur.label(cur.TRIMESTER_LABELS, lang, spec.trimester)}"
        paper.heading = heading

    if spec.school_name:
        paper.left_rows.append((paper.t("school"), spec.school_name))
    if spec.teacher_name:
        paper.left_rows.append((paper.t("teacher"), spec.teacher_name))
    if spec.school_year:
        paper.left_rows.append((paper.t("school_year"), spec.school_year))

    if spec.subject:
        paper.right_rows.append((paper.t("subject"), spec.subject))
    if spec.level:
        level = cur.label(cur.LEVEL_LABELS, lang, spec.level, spec.level)
        if spec.section:
            level = f"{level} — {cur.label(cur.SECTION_LABELS, lang, spec.section, spec.section)}"
        paper.right_rows.append((paper.t("level"), level))
    if spec.duration_minutes:
        paper.right_rows.append(
            (paper.t("duration"), f"{spec.duration_minutes} {paper.t('minutes')}")
        )
    return paper


def _exercise_heading(paper: _Paper, index: int, ex: Exercise) -> str:
    ex_label = paper.t("exercise")
    sep = "" if ex_label.endswith("°") else " "
    heading = f"{ex_label}{sep}{index}"
    if ex.type in ("mcq", "true_false"):
        heading = f"{heading} — {paper.t('type_' + ex.type)}"
    return f"{heading} ({ex.points} {paper.t('points')})"


def _choices(paper: _Paper, ex: Exercise) -> list[str]:
    if ex.choices:
        return list(ex.choices)
    if ex.type == "true_false":
        return [paper.t("true"), paper.t("false")]
    return []


def _student_strip(paper: _Paper, total_points: int) -> str:
    return (
        f"{paper.t('student_name')} : {'.' * 34}   "
        f"{paper.t('class')} : {'.' * 10}   "
        f"{paper.t('grade')} : {'.' * 8} / {total_points}"
    )


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #


def _pdf_text(value: str, *, keep_linebreaks: bool = False) -> str:
    """Escape user/LLM-provided text for ReportLab's mini-markup.

    ReportLab ``Paragraph`` interprets a subset of HTML-like tags, so raw ``<``,
    ``>``, and ``&`` from exam content could break rendering or inject markup.
    Escape first, then (optionally) re-introduce ``<br/>`` for real newlines.
    """
    escaped = escape(value or "")
    if keep_linebreaks:
        escaped = escaped.replace("\n", "<br/>")
    return escaped


def _rtl_lines(text: str, font_name: str, font_size: float, max_width: float) -> list[str]:
    """Wrap logical-order Arabic text into visual lines that fit ``max_width``.

    Each returned line is already shaped and bidi-reordered, ready to be drawn
    left-to-right by ReportLab.
    """
    out: list[str] = []
    for logical in (text or "").split("\n"):
        reshaped = arabic_reshaper.reshape(logical)
        current = ""
        for word in reshaped.split(" "):
            candidate = word if not current else f"{current} {word}"
            if not current or pdfmetrics.stringWidth(candidate, font_name, font_size) <= max_width:
                current = candidate
            else:
                out.append(current)
                current = word
        out.append(current)
    return [get_display(line) for line in out]


def _para(
    text: str,
    style: ParagraphStyle,
    paper: _Paper,
    width: float,
    *,
    bold: bool = False,
    italic: bool = False,
) -> Paragraph:
    """Build a Paragraph, handling Arabic shaping/RTL when the paper is RTL."""
    if paper.rtl:
        # Width is reduced slightly: stringWidth ignores kerning/justification.
        lines = _rtl_lines(text, style.fontName, style.fontSize, width * 0.97)
        return Paragraph("<br/>".join(escape(l) for l in lines), style)
    markup = _pdf_text(text, keep_linebreaks=True)
    if bold:
        markup = f"<b>{markup}</b>"
    if italic:
        markup = f"<i>{markup}</i>"
    return Paragraph(markup, style)


def _pdf_styles(paper: _Paper) -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    if paper.rtl and _ensure_arabic_fonts():
        align = TA_RIGHT
        body = ParagraphStyle(
            "ArBody", fontName=ARABIC_FONT, fontSize=12, leading=20, alignment=align
        )
        return {
            "title": ParagraphStyle(
                "ArTitle", fontName=ARABIC_FONT_BOLD, fontSize=17, leading=26, alignment=TA_CENTER
            ),
            "subtitle": ParagraphStyle(
                "ArSubtitle", fontName=ARABIC_FONT, fontSize=13, leading=20, alignment=TA_CENTER
            ),
            "h2": ParagraphStyle(
                "ArH2",
                fontName=ARABIC_FONT_BOLD,
                fontSize=13,
                leading=22,
                alignment=align,
                spaceBefore=8,
                spaceAfter=4,
            ),
            "body": body,
            "bold": ParagraphStyle("ArBold", parent=body, fontName=ARABIC_FONT_BOLD),
            "small": ParagraphStyle("ArSmall", parent=body, fontSize=11, leading=17),
        }
    if paper.rtl:
        # Fonts missing: degrade to LTR Helvetica rather than failing the export.
        paper.rtl = False
    body = base["BodyText"]
    return {
        "title": base["Title"],
        "subtitle": ParagraphStyle(
            "Subtitle", parent=base["Heading3"], alignment=TA_CENTER, spaceAfter=6
        ),
        "h2": ParagraphStyle("ExerciseHeader", parent=base["Heading2"], spaceAfter=6),
        "body": body,
        "bold": body,
        "small": ParagraphStyle("Small", parent=body, fontSize=9.5, leading=12),
    }


def _pdf_header_table(paper: _Paper, styles: dict, width: float) -> Optional[Table]:
    if not paper.has_header:
        return None
    col_w = width / 2
    n = max(len(paper.left_rows), len(paper.right_rows))
    left = list(paper.left_rows) + [("", "")] * (n - len(paper.left_rows))
    right = list(paper.right_rows) + [("", "")] * (n - len(paper.right_rows))

    def cell(pair: tuple[str, str]) -> Paragraph:
        label_, value = pair
        text = f"{label_} : {value}" if label_ else ""
        return _para(text, styles["small"], paper, col_w - 0.6 * cm)

    if paper.rtl:
        # Mirror the columns so the "administrative" block sits on the right.
        rows = [[cell(r), cell(l)] for l, r in zip(left, right)]
    else:
        rows = [[cell(l), cell(r)] for l, r in zip(left, right)]

    table = Table(rows, colWidths=[col_w, col_w])
    table.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.8, colors.black),
                ("LINEBEFORE", (1, 0), (1, -1), 0.8, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def to_pdf(exam: ExamContent, spec: Optional[ExamSpec] = None) -> bytes:
    paper = _paper_from_spec(spec)
    styles = _pdf_styles(paper)

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=exam.title or paper.heading or "Exam",
    )
    width = pdf.width
    story: list = []

    header = _pdf_header_table(paper, styles, width)
    if header is not None:
        story.extend([header, Spacer(1, 0.4 * cm)])

    if paper.heading:
        story.append(_para(paper.heading, styles["title"], paper, width))
        if exam.title:
            story.append(_para(exam.title, styles["subtitle"], paper, width))
    else:
        story.append(_para(exam.title or "Exam", styles["title"], paper, width))

    if paper.has_header or paper.heading:
        story.append(Spacer(1, 0.2 * cm))
        story.append(_para(_student_strip(paper, exam.total_points), styles["small"], paper, width))
    else:
        story.append(
            _para(
                f"{paper.t('total')} : {exam.total_points} {paper.t('points')}",
                styles["body"],
                paper,
                width,
                bold=True,
            )
        )
    story.append(Spacer(1, 0.5 * cm))

    for i, ex in enumerate(exam.exercises, start=1):
        story.append(_para(_exercise_heading(paper, i, ex), styles["h2"], paper, width))
        story.append(_para(ex.question, styles["body"], paper, width))
        choices = _choices(paper, ex)
        if choices:
            if paper.rtl:
                # ListFlowable puts bullets on the left; for RTL prefix a marker
                # (U+2022 is present in Amiri) so it lands on the right after bidi.
                for c in choices:
                    story.append(_para(f"• {c}", styles["body"], paper, width - 0.5 * cm))
            else:
                story.append(
                    ListFlowable(
                        [ListItem(_para(c, styles["body"], paper, width)) for c in choices],
                        bulletType="bullet",
                    )
                )
        story.append(Spacer(1, 0.3 * cm))

    story.append(PageBreak())
    story.append(_para(paper.t("answer_key"), styles["title"], paper, width))
    story.append(Spacer(1, 0.3 * cm))
    for i, ex in enumerate(exam.exercises, start=1):
        story.append(_para(_exercise_heading(paper, i, ex), styles["bold"], paper, width, bold=True))
        story.append(_para(ex.answer, styles["body"], paper, width))
        if ex.explanation:
            story.append(
                _para(
                    f"{paper.t('explanation')} : {ex.explanation}",
                    styles["body"],
                    paper,
                    width,
                    italic=True,
                )
            )
        story.append(Spacer(1, 0.2 * cm))

    pdf.build(story)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #


def _docx_rtl(paragraph, *, center: bool = False) -> None:
    """Mark a python-docx paragraph (and its runs) as right-to-left."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:bidi"))
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def _docx_paragraph(doc_or_cell, text: str, paper: _Paper, *, style=None, bold=False, italic=False, center=False):
    p = doc_or_cell.add_paragraph(style=style) if style else doc_or_cell.add_paragraph()
    run = p.add_run(text or "")
    run.bold = bold or None
    run.italic = italic or None
    if paper.rtl:
        _docx_rtl(p, center=center)
    elif center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _docx_heading(doc, text: str, paper: _Paper, level: int, *, center: bool = False):
    h = doc.add_heading(text or "", level=level)
    if paper.rtl:
        _docx_rtl(h, center=center)
    elif center:
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return h


def _docx_header_table(doc, paper: _Paper) -> None:
    if not paper.has_header:
        return
    n = max(len(paper.left_rows), len(paper.right_rows))
    left = list(paper.left_rows) + [("", "")] * (n - len(paper.left_rows))
    right = list(paper.right_rows) + [("", "")] * (n - len(paper.right_rows))
    table = doc.add_table(rows=n, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if paper.rtl:
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    for row_idx in range(n):
        for col_idx, (label_, value) in enumerate((left[row_idx], right[row_idx])):
            cell = table.cell(row_idx, col_idx)
            cell.paragraphs[0].text = ""
            text = f"{label_} : {value}" if label_ else ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(10)
            if paper.rtl:
                _docx_rtl(p)
    doc.add_paragraph()


def to_docx(exam: ExamContent, spec: Optional[ExamSpec] = None) -> bytes:
    paper = _paper_from_spec(spec)
    doc = Document()

    _docx_header_table(doc, paper)

    if paper.heading:
        _docx_heading(doc, paper.heading, paper, level=1, center=True)
        if exam.title:
            _docx_paragraph(doc, exam.title, paper, bold=True, center=True)
    else:
        _docx_heading(doc, exam.title or "Exam", paper, level=1, center=True)

    if paper.has_header or paper.heading:
        _docx_paragraph(doc, _student_strip(paper, exam.total_points), paper)
    else:
        _docx_paragraph(
            doc,
            f"{paper.t('total')} : {exam.total_points} {paper.t('points')}",
            paper,
            bold=True,
            center=True,
        )

    for i, ex in enumerate(exam.exercises, start=1):
        h = _docx_heading(doc, _exercise_heading(paper, i, ex), paper, level=2)
        for r in h.runs:
            r.font.size = Pt(13)
        _docx_paragraph(doc, ex.question, paper)
        for c in _choices(paper, ex):
            _docx_paragraph(doc, c, paper, style="List Bullet")

    doc.add_page_break()
    _docx_heading(doc, paper.t("answer_key"), paper, level=1, center=True)
    for i, ex in enumerate(exam.exercises, start=1):
        para = doc.add_paragraph()
        para.add_run(f"{_exercise_heading(paper, i, ex)} : ").bold = True
        para.add_run(ex.answer)
        if paper.rtl:
            _docx_rtl(para)
        if ex.explanation:
            ep = doc.add_paragraph()
            ep.add_run(f"{paper.t('explanation')} : ").italic = True
            ep.add_run(ex.explanation)
            if paper.rtl:
                _docx_rtl(ep)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
